"""
Generate draft analytical report as DOCX.
"""
import os
from datetime import datetime
from decimal import Decimal
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from app import db
from app.models.case import Case
from app.models.invoice import Invoice, InvoiceLine
from app.models.match import ProductMatch
from app.models.flag import AnalyticalFlag
from app.models.report import ReportVersion
from flask import current_app
from sqlalchemy import func


def generate_draft_report(case_id: int, user_id: int) -> ReportVersion:
    case = Case.query.get_or_404(case_id)

    # Aggregate stats
    purchase_sum = (
        db.session.query(func.coalesce(func.sum(InvoiceLine.amount), 0))
        .join(Invoice)
        .filter(InvoiceLine.case_id == case_id, Invoice.operation_type == "purchase")
        .scalar()
    )
    sale_sum = (
        db.session.query(func.coalesce(func.sum(InvoiceLine.amount), 0))
        .join(Invoice)
        .filter(InvoiceLine.case_id == case_id, Invoice.operation_type == "sale")
        .scalar()
    )
    purchase_count = (
        InvoiceLine.query.join(Invoice)
        .filter(InvoiceLine.case_id == case_id, Invoice.operation_type == "purchase")
        .count()
    )
    sale_count = (
        InvoiceLine.query.join(Invoice)
        .filter(InvoiceLine.case_id == case_id, Invoice.operation_type == "sale")
        .count()
    )
    confirmed_matches = ProductMatch.query.filter_by(case_id=case_id, status="confirmed").count()
    flags = AnalyticalFlag.query.filter_by(case_id=case_id).all()

    # Build structured content
    content = {
        "header": {
            "title": f"Аналитическая справка по делу {case.number}",
            "organization": case.organization_name,
            "bin": case.bin,
            "period": f"{case.period_start} — {case.period_end}" if case.period_start else "",
            "basis": case.basis,
        },
        "summary": {
            "purchase_sum": float(purchase_sum) if purchase_sum else 0,
            "sale_sum": float(sale_sum) if sale_sum else 0,
            "purchase_count": purchase_count,
            "sale_count": sale_count,
            "confirmed_matches": confirmed_matches,
            "flags_count": len(flags),
        },
        "flags": [f.to_dict() for f in flags if f.status in ("new", "confirmed", "under_review")],
        "generated_at": datetime.utcnow().isoformat(),
    }

    # Create DOCX
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("АНАЛИТИЧЕСКАЯ СПРАВКА")
    run.bold = True
    run.font.size = Pt(16)

    doc.add_paragraph(f"Дело № {case.number}")
    doc.add_paragraph(f"Организация: {case.organization_name or '—'} (БИН: {case.bin or '—'})")
    if case.period_start:
        doc.add_paragraph(f"Период анализа: {case.period_start.strftime('%d.%m.%Y')} — {case.period_end.strftime('%d.%m.%Y') if case.period_end else '—'}")
    if case.basis:
        doc.add_paragraph(f"Основание: {case.basis}")

    doc.add_heading("1. Общие сведения о приобретении и реализации", level=1)
    p = doc.add_paragraph()
    p.add_run(
        f"Согласно обработанным сведениям электронных счетов-фактур за указанный период, "
        f"сумма приобретения по исследуемой организации составила {float(purchase_sum):,.2f} тенге "
        f"({purchase_count} товарных позиций), а сумма реализации — {float(sale_sum):,.2f} тенге "
        f"({sale_count} товарных позиций)."
    )
    p = doc.add_paragraph()
    p.add_run(
        "Указанные значения отражают совокупные обороты по доступной выборке документов и сами по себе "
        "не подтверждают отсутствие иных источников приобретения или наличие нарушения. "
        "Для уточнения обстоятельств необходимо проверить состав товарных позиций, остатки на начало периода, "
        "операции за предыдущие периоды, возвраты и другие доступные сведения."
    )

    doc.add_heading("2. Сопоставление товарных позиций", level=1)
    doc.add_paragraph(f"Подтверждённых сопоставлений: {confirmed_matches}.")

    doc.add_heading("3. Аналитические обстоятельства", level=1)
    if flags:
        for f in flags:
            if f.status in ("new", "confirmed", "under_review"):
                doc.add_paragraph(f"• [{f.category}] {f.title}", style="List Bullet")
                if f.description:
                    doc.add_paragraph(f"  {f.description}")
    else:
        doc.add_paragraph("Существенных аналитических сигналов не выявлено (по текущим правилам).")

    doc.add_heading("4. Выводы", level=1)
    doc.add_paragraph(
        "[Раздел заполняется аналитиком вручную. Система не формирует юридически значимые выводы автоматически.]"
    )

    doc.add_heading("5. Источники", level=1)
    doc.add_paragraph("• Загруженные выгрузки ЭСФ / Excel / CSV")
    doc.add_paragraph("• Результаты автоматической обработки и ручной проверки сотрудником")

    doc.add_paragraph()
    footer = doc.add_paragraph()
    footer.add_run(f"Черновик сформирован системой Analytics Desk {datetime.utcnow().strftime('%d.%m.%Y %H:%M')} UTC").italic = True

    # Save file
    upload_root = current_app.config["UPLOAD_FOLDER"]
    reports_dir = os.path.join(upload_root, "reports", str(case_id))
    os.makedirs(reports_dir, exist_ok=True)
    version = ReportVersion.query.filter_by(case_id=case_id).count() + 1
    filename = f"report_v{version}_{case.number.replace(' ', '_')}.docx"
    filepath = os.path.join(reports_dir, filename)
    doc.save(filepath)

    rv = ReportVersion(
        case_id=case_id,
        version=version,
        title=f"Аналитическая справка — {case.number} (v{version})",
        status="draft",
        content_json=content,
        file_path=filepath,
        created_by_id=user_id,
    )
    db.session.add(rv)
    db.session.commit()
    return rv
